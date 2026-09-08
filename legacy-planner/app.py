import cgi
import copy
import json
import os
import re
import subprocess
import sys
import tempfile
from http.server import HTTPServer, SimpleHTTPRequestHandler
from pathlib import Path

from furnace_domain import FURNACES, LONG_THRESHOLD, LoadingEngine, RING_HEIGHT, RING_LIMIT

def runtime_root():
    configured = os.environ.get('PREHEAT_APP_ROOT')
    if configured:
        return Path(configured).resolve()
    if getattr(sys, 'frozen', False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


ROOT = runtime_root()
DATA_ROOT = Path(os.environ.get(
    'PREHEAT_DATA_ROOT',
    Path(os.environ.get('LOCALAPPDATA', Path.home())) / 'PreheatFurnaceWorkbench',
)).resolve()
DATA_ROOT.mkdir(parents=True, exist_ok=True)
LENGTHS_FILE = DATA_ROOT / 'billet_lengths.json'
DEFAULT_LENGTHS_FILE = ROOT / 'assets' / 'billet_lengths.json'
if not LENGTHS_FILE.exists() and DEFAULT_LENGTHS_FILE.exists():
    LENGTHS_FILE.write_bytes(DEFAULT_LENGTHS_FILE.read_bytes())
NODE_EXE = Path(os.environ.get('PREHEAT_NODE_EXE', ROOT / 'runtime' / 'node' / 'node.exe'))
EXPORT_SCRIPT = ROOT / 'export_loading_map.mjs'
EXPORT_TEMPLATE = Path(os.environ.get(
    'PREHEAT_EXPORT_TEMPLATE', ROOT / 'assets' / '锻造预热炉装炉图 v5.xlsx',
))
VIRTUAL_FURNACE = 'VIRTUAL'
VIRTUAL_FURNACE_Z = 'VIRTUAL_Z'


def number(value):
    try:
        return int(float(str(value or '').strip() or 0))
    except (TypeError, ValueError):
        return 0


def parse_forging_numbers(value, quantity):
    """Expand the plan's forging-number field only when it exactly covers the planned quantity."""
    text = str(value or '').strip()
    if not text or quantity <= 0:
        return None
    text = (text.replace('，', ',').replace('、', ',').replace('；', ',')
                .replace('~', '~').replace('～', '~'))
    numbers = []
    token_pattern = r'\[?\s*([A-Za-z]\d+(?:-[A-Za-z0-9]+)*(?:~[A-Za-z0-9]+)?)\s*\]?(?:\s*(\d+)\s*件)?'
    matches = list(re.finditer(token_pattern, text))
    if not matches or re.sub(r'[\s,\[\]]', '', re.sub(token_pattern, '', text)):
        return None
    for match in matches:
        token = match.groups()
        start, repeat = token
        range_match = re.fullmatch(r'(.+-)([A-Za-z]+|\d+)~([A-Za-z]+|\d+)', start)
        if range_match:
            prefix, first, last = range_match.groups()
            if first.isdigit() and last.isdigit():
                expanded = [f'{prefix}{index}' for index in range(int(first), int(last) + 1)]
            elif len(first) == len(last) == 1 and first.isalpha() and last.isalpha():
                first, last = first.upper(), last.upper()
                expanded = [f'{prefix}{chr(index)}' for index in range(ord(first), ord(last) + 1)]
            else:
                return None
            if not expanded:
                return None
            numbers.extend(expanded)
        else:
            numbers.extend([start] * (number(repeat) or 1))
    return numbers if len(numbers) == quantity else None


def read_lengths():
    if not LENGTHS_FILE.exists():
        return {}
    try:
        data = json.loads(LENGTHS_FILE.read_text(encoding='utf-8'))
        return {str(k).strip(): number(v) for k, v in data.items() if number(v) > 0}
    except (OSError, json.JSONDecodeError):
        return {}


def save_lengths(lengths):
    LENGTHS_FILE.write_text(json.dumps(lengths, ensure_ascii=False, indent=2), encoding='utf-8')


def length_map(path):
    from collections import Counter

    import openpyxl

    ws = openpyxl.load_workbook(path, data_only=True, read_only=True).active
    headers = {str(value or '').strip(): index for index, value in enumerate(next(ws.iter_rows(values_only=True)), start=1)}
    required = {'产品编号', '长度'}
    if not required.issubset(headers):
        raise ValueError('棒料登记表缺少“产品编号”或“长度”列')
    code_index, length_index = headers['产品编号'] - 1, headers['长度'] - 1
    values_by_code = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        code = str(row[code_index] or '').strip() if code_index < len(row) else ''
        length = number(row[length_index]) if length_index < len(row) else 0
        if code and length > 0:
            values_by_code.setdefault(code, []).append(length)
    return {code: Counter(values).most_common(1)[0][0] for code, values in values_by_code.items()}


def load_items(plan_path, lengths):
    import openpyxl

    ws = openpyxl.load_workbook(plan_path, data_only=True).active
    items = []
    for row in range(2, ws.max_row + 1):
        code = str(ws[f'B{row}'].value or '').strip()
        alloy = str(ws[f'C{row}'].value or '').strip()
        forging_numbers = parse_forging_numbers(ws[f'D{row}'].value, number(ws[f'E{row}'].value))
        quantity = number(ws[f'E{row}'].value)
        if not code or quantity <= 0 or '外发' in alloy or '下料' in alloy:
            continue
        is_ring = any(token in alloy for token in ('锻造碾环', '碾环', '冲孔'))
        height = RING_HEIGHT if is_ring else lengths.get(code, 0)
        if height <= 0:
            items.append({'code': code, 'alloy': alloy, 'quantity': quantity, 'height_mm': 0,
                          'diameter_mm': 630, 'is_ring': False, 'invalid': '未记住该棒料长度'})
            continue
        base = {'code': code, 'alloy': alloy, 'height_mm': height,
                'diameter_mm': 950 if '950' in alloy else 630, 'is_ring': is_ring}
        if forging_numbers:
            items.extend({**base, 'quantity': 1, 'forging_no': forging_no}
                         for forging_no in forging_numbers)
        else:
            items.append({**base, 'quantity': quantity})
    missing = [x for x in items if x.get('invalid')]
    valid = [x for x in items if not x.get('invalid')]
    return valid, [{'code': x['code'], 'forging_no': x.get('forging_no'),
                    'reason': x['invalid'], 'quantity': x['quantity']} for x in missing]


def refresh_column(column):
    column['items'].sort(key=lambda item: -item['height_mm'])
    height = 0
    for item in column['items']:
        height += item['height_mm']
        item['cumulative_mm'] = height
    column['height_mm'] = height


def refresh_virtual_summary(result):
    virtual_items = result['furnaces'][VIRTUAL_FURNACE]['columns'][0]['items']
    result['unassigned'] = [
        {'code': item['code'], 'forging_no': item.get('forging_no'),
         'reason': item.get('reason', '未分配')}
        for item in virtual_items
    ]
    result['total_unassigned'] = len(virtual_items)
    result['total_placed'] = sum(
        furnace['total'] for furnace_id, furnace in result['furnaces'].items()
        if furnace_id not in (VIRTUAL_FURNACE, VIRTUAL_FURNACE_Z)
    )


def add_virtual_furnace(result, source_items):
    """Expose every not-yet-loaded unit as draggable inventory."""
    by_identity = {(item['code'], item.get('forging_no')): item for item in source_items}
    virtual_items = []
    for entry in result.get('unassigned', []):
        source = by_identity.get((entry['code'], entry.get('forging_no')), {})
        quantity = number(entry.get('quantity')) or 1
        for _ in range(quantity):
            virtual_items.append({
                'code': entry['code'], 'forging_no': entry.get('forging_no'),
                'alloy': source.get('alloy', ''),
                'height_mm': source.get('height_mm', 0), 'cumulative_mm': 0,
                'type': '待分配', 'is_ring': source.get('is_ring', False),
                'reason': entry.get('reason', '未分配'),
            })
    result['furnaces'][VIRTUAL_FURNACE] = {
        'spec': {'furnace_id': VIRTUAL_FURNACE, 'slots': 0, 'height_mm': 0,
                 'speed': '', 'coefficient': 0, 'base_heat_h': 0,
                 'role': '未分配暂存（最低优先级）', 'is_large': False,
                 'is_ring': False, 'ring_slots': 0, 'is_virtual': True},
        'total': len(virtual_items), 'columns': [{'height_mm': 0, 'items': virtual_items}],
        'max_height_mm': 0,
    }
    result['furnaces'][VIRTUAL_FURNACE_Z] = {
        'spec': {'furnace_id': VIRTUAL_FURNACE_Z, 'slots': 0, 'height_mm': 0,
                 'speed': '', 'coefficient': 0, 'base_heat_h': 0,
                 'role': '手动暂存，不参与自动排炉', 'is_large': False,
                 'is_ring': False, 'ring_slots': 0, 'is_virtual': True,
                 'manual_only': True},
        'total': 0, 'columns': [{'height_mm': 0, 'items': []}],
        'max_height_mm': 0,
    }
    refresh_virtual_summary(result)


def move_item(result, source_furnace, source_column, item_index, target_furnace, target_column=None):
    furnaces = result['furnaces']
    if source_furnace not in furnaces or target_furnace not in furnaces:
        raise ValueError('炉号不存在')
    source_columns = furnaces[source_furnace]['columns']
    if not 0 <= source_column < len(source_columns):
        raise ValueError('来源列不存在')
    source = source_columns[source_column]
    if not 0 <= item_index < len(source['items']):
        raise ValueError('来源棒料不存在')
    item = source['items'][item_index]
    target = furnaces[target_furnace]
    if target_furnace in (VIRTUAL_FURNACE, VIRTUAL_FURNACE_Z):
        source['items'].pop(item_index)
        refresh_column(source)
        target['columns'][0]['items'].append(item)
        target['total'] += 1
        for furnace in furnaces.values():
            furnace['total'] = sum(len(col['items']) for col in furnace['columns'])
            furnace['max_height_mm'] = max((col['height_mm'] for col in furnace['columns']), default=0)
        refresh_virtual_summary(result)
        return result
    if item['height_mm'] <= 0:
        raise ValueError('请先补充该物料的棒料长度，再拖入实体炉')
    target_spec = target['spec']
    # This route is only used by manual drag/drop and batch adjustments.
    # The automatic allocator keeps the 1408 alloy lock in LoadingEngine.
    limit = RING_LIMIT if item.get('is_ring') and target_spec['is_ring'] else target_spec['height_mm']
    target_columns = target['columns']
    if target_column is not None:
        target_index = number(target_column)
        if not 0 <= target_index < len(target_columns):
            raise ValueError(f'{target_furnace} 第 {target_index + 1} 列不存在')
        if target_columns[target_index]['height_mm'] + item['height_mm'] > limit:
            raise ValueError(f'{target_furnace} 第 {target_index + 1} 列高度不足')
    else:
        target_index = None
    candidates = [(column['height_mm'], index) for index, column in enumerate(target_columns)
                  if column['height_mm'] + item['height_mm'] <= limit]
    if target_index is None and not candidates:
        raise ValueError(f'{target_furnace} 没有满足高度约束的空位')
    if target_index is None:
        _, target_index = min(candidates)
    source['items'].pop(item_index)
    refresh_column(source)
    target_columns[target_index]['items'].append(item)
    refresh_column(target_columns[target_index])
    for furnace in furnaces.values():
        furnace['total'] = sum(len(col['items']) for col in furnace['columns'])
        furnace['max_height_mm'] = max((col['height_mm'] for col in furnace['columns']), default=0)
    refresh_virtual_summary(result)
    return result


class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(ROOT / 'static'), **kwargs)

    def _json(self, data, status=200):
        self.send_response(status)
        self.send_header('Content-Type', 'application/json; charset=utf-8')
        self.end_headers()
        self.wfile.write(json.dumps(data, ensure_ascii=False).encode())

    def _xlsx(self, path):
        self.send_response(200)
        self.send_header('Content-Type', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.send_header('Content-Disposition', 'attachment; filename="furnace-loading-map.xlsx"')
        self.send_header('Content-Length', str(path.stat().st_size))
        self.end_headers()
        self.wfile.write(path.read_bytes())

    def _read_json(self):
        length = number(self.headers.get('Content-Length'))
        return json.loads(self.rfile.read(length).decode('utf-8'))

    def do_GET(self):
        if self.path == '/':
            self.path = '/index.html'
            return super().do_GET()
        if self.path == '/api/config':
            return self._json({'furnaces': [f.__dict__ for f in FURNACES], 'ring_height_mm': RING_HEIGHT,
                               'long_threshold_mm': LONG_THRESHOLD})
        if self.path == '/api/billet-lengths':
            return self._json({'lengths': read_lengths()})
        return super().do_GET()

    def do_POST(self):
        try:
            if self.path == '/api/loading-plans/export':
                body = self._read_json()
                with tempfile.TemporaryDirectory() as temp:
                    input_path = Path(temp) / 'loading-result.json'
                    output_path = Path(temp) / 'furnace-loading-map.xlsx'
                    input_path.write_text(json.dumps(body['result'], ensure_ascii=False), encoding='utf-8')
                    if not EXPORT_TEMPLATE.exists():
                        raise ValueError('未找到装炉图导出模板')
                    process = subprocess.run(
                        [str(NODE_EXE), str(EXPORT_SCRIPT), str(input_path), str(output_path), str(EXPORT_TEMPLATE)],
                        capture_output=True, text=True, timeout=60, check=False,
                    )
                    if process.returncode != 0 or not output_path.exists():
                        raise ValueError(f'装炉图导出失败：{process.stderr.strip() or process.stdout.strip()}')
                    return self._xlsx(output_path)
            if self.path == '/api/billet-lengths':
                body = self._read_json()
                code, length = str(body.get('code', '')).strip(), number(body.get('length_mm'))
                if not code or length <= 0:
                    raise ValueError('请输入物料编码和大于 0 的长度')
                lengths = read_lengths()
                lengths[code] = length
                save_lengths(lengths)
                return self._json({'ok': True, 'code': code, 'length_mm': length})
            if self.path == '/api/loading-plans/move':
                body = self._read_json()
                result = move_item(body['result'], body['source_furnace'], number(body['source_column']),
                                   number(body['item_index']), body['target_furnace'], body.get('target_column'))
                return self._json({'ok': True, 'result': result})
            if self.path == '/api/loading-plans/batch-move':
                body = self._read_json()
                moves = body.get('moves', [])
                target = body.get('target_furnace')
                if not moves or not target:
                    raise ValueError('请选择至少一根棒料和目标炉')
                result = copy.deepcopy(body['result'])
                # Process each source column from its last item upward so indices remain stable.
                normalized = [{
                    'source_furnace': move.get('source_furnace', move.get('furnace')),
                    'source_column': move.get('source_column', move.get('column')),
                    'item_index': move.get('item_index', move.get('item')),
                } for move in moves]
                if any(not move['source_furnace'] for move in normalized):
                    raise ValueError('批量换炉缺少来源炉信息')
                ordered = sorted(normalized, key=lambda move: (
                    move['source_furnace'], number(move['source_column']), -number(move['item_index'])
                ))
                for move in ordered:
                    move_item(result, move['source_furnace'], number(move['source_column']),
                              number(move['item_index']), target)
                return self._json({'ok': True, 'result': result, 'moved': len(moves)})
            if self.path != '/api/loading-plans':
                return self._json({'ok': False, 'message': '接口不存在'}, 404)
            form = cgi.FieldStorage(fp=self.rfile, headers=self.headers,
                                    environ={'REQUEST_METHOD': 'POST', 'CONTENT_TYPE': self.headers['Content-Type']})
            plan = form['plan_file']
            billet = form['billet_file'] if 'billet_file' in form else None
            if not plan.filename:
                raise ValueError('请选择当天生产计划')
            lengths = read_lengths()
            with tempfile.TemporaryDirectory() as temp:
                plan_path = Path(temp) / 'plan.xlsx'
                plan_path.write_bytes(plan.file.read())
                if billet is not None and bool(billet.filename):
                    billet_path = Path(temp) / 'billet.xlsx'
                    billet_path.write_bytes(billet.file.read())
                    lengths.update(length_map(billet_path))
                    save_lengths(lengths)
                items, missing = load_items(plan_path, lengths)
            result = LoadingEngine().allocate(items)
            result['unassigned'].extend(missing)
            result['total_unassigned'] = len(result['unassigned'])
            add_virtual_furnace(result, items)
            return self._json({'ok': True, 'items': items, 'result': result,
                               'remembered_lengths': len(lengths)})
        except Exception as exc:
            return self._json({'ok': False, 'message': f'导入或排炉失败：{exc}'}, 422)


if __name__ == '__main__':
    print('http://127.0.0.1:8760')
    HTTPServer(('127.0.0.1', 8760), Handler).serve_forever()
