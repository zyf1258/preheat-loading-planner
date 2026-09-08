import json
import tempfile
import threading
import unittest
from io import BytesIO
from pathlib import Path
from urllib.error import HTTPError
from urllib.request import Request, urlopen

import app
import legacy_adapter
import openpyxl

from preheat_planner.service import DEFAULT_SERVICE


def _detail_result(*furnaces):
    return {
        "furnaces": {
            furnace_id: {
                "spec": {"is_virtual": False, "is_ring": is_ring},
                "columns": [{"items": items}],
            }
            for furnace_id, is_ring, items in furnaces
        }
    }


def _item(code, forging_no, *, is_ring, quantity=1, height_mm=430):
    return {
        "code": code,
        "forging_no": forging_no,
        "is_ring": is_ring,
        "quantity": quantity,
        "height_mm": height_mm,
    }


class LoadingMapExportTest(unittest.TestCase):
    def test_ring_furnace_merges_same_material_and_consecutive_forging_numbers(self):
        result = _detail_result(
            ("1404", True, [
                _item("F2.6061.0322", "D28945-A-2", is_ring=True),
                _item("F2.6061.0322", "D28945-B-1", is_ring=True),
                _item("F2.6061.0322", "D28945-B-2", is_ring=True),
                _item("F2.6061.0375", "D28961-A-1", is_ring=True),
            ]),
        )

        self.assertEqual(
            legacy_adapter._export_furnace_details(result),
            [
                ["1404", "F2.6061.0322", "D28945-A-2", 1, "430mm", None],
                ["1404", "F2.6061.0322", "D28945-B-1~2", 2, "430mm", None],
                ["1404", "F2.6061.0375", "D28961-A-1", 1, "430mm", None],
            ],
        )

    def test_nonconsecutive_ring_numbers_remain_separate(self):
        result = _detail_result(
            ("1406", True, [
                _item("F2.6061.0322", "D28945-B-1", is_ring=True),
                _item("F2.6061.0322", "D28945-B-3", is_ring=True),
            ]),
        )

        self.assertEqual(
            [row[2] for row in legacy_adapter._export_furnace_details(result)],
            ["D28945-B-1", "D28945-B-3"],
        )

    def test_non_ring_estimated_weight_uses_length_times_0791(self):
        result = _detail_result(
            ("1405", False, [
                _item("F3.6061.0201", "D30001-A", is_ring=False, height_mm=703),
            ]),
        )

        self.assertEqual(
            legacy_adapter._export_furnace_details(result),
            [["1405", "F3.6061.0201", "D30001-A", 1, "703mm", 556.07]],
        )

    def test_forging_furnace_does_not_merge_same_material(self):
        result = _detail_result(
            ("1405", False, [
                _item("F2.6061.0322", "D28945-B-1", is_ring=True),
                _item("F2.6061.0322", "D28945-B-2", is_ring=True),
            ]),
        )

        self.assertEqual(
            legacy_adapter._export_furnace_details(result),
            [
                ["1405", "F2.6061.0322", "D28945-B-1", 1, "430mm", None],
                ["1405", "F2.6061.0322", "D28945-B-2", 1, "430mm", None],
            ],
        )

    def test_ring_material_cells_are_merged_in_export_workbook(self):
        result = _detail_result(
            ("1404", True, [
                _item("F2.6061.0322", "D28945-A-1", is_ring=True),
                _item("F2.6061.0322", "D28945-A-2", is_ring=True),
                _item("F2.6061.0322", "D28945-B-1", is_ring=True),
                _item("F2.6061.0322", "D28945-B-2", is_ring=True),
                _item("F2.6061.0375", "D28961-A-1", is_ring=True),
            ]),
        )
        template = Path(__file__).parents[1] / "legacy-planner" / "assets" / "锻造预热炉装炉图 v5.xlsx"

        workbook_bytes = legacy_adapter.export_loading_map_with_openpyxl(result, template)
        sheet = openpyxl.load_workbook(BytesIO(workbook_bytes))["预热炉装炉图"]

        merged_ranges = {str(value) for value in sheet.merged_cells.ranges}
        self.assertNotIn("N10:N12", merged_ranges)
        self.assertIn("O10:O11", {str(value) for value in sheet.merged_cells.ranges})
        self.assertEqual(sheet["P10"].value, "D28945-A-1~2")
        self.assertEqual(sheet["P11"].value, "D28945-B-1~2")
        self.assertEqual(sheet["Q11"].value, 2)
        self.assertEqual(sheet["O11"].value, None)
        self.assertEqual(sheet["S10"].value, None)
        self.assertEqual(sheet["S11"].value, None)

    def test_export_writes_non_ring_estimated_weight_column(self):
        result = _detail_result(
            ("1405", False, [
                _item("F3.6061.0201", "D30001-A", is_ring=False, height_mm=703),
            ]),
        )
        template = Path(__file__).parents[1] / "legacy-planner" / "assets" / "锻造预热炉装炉图 v5.xlsx"

        workbook_bytes = legacy_adapter.export_loading_map_with_openpyxl(result, template)
        sheet = openpyxl.load_workbook(BytesIO(workbook_bytes))["预热炉装炉图"]

        self.assertEqual(sheet["S9"].value, "预估重量")
        self.assertEqual(sheet["S10"].value, 556.07)

    def test_import_exported_map_restores_items_and_can_export_again(self):
        result = _detail_result(
            ("1404", True, [
                _item("F2.6061.0322", "D28945-B-1", is_ring=True),
                _item("F2.6061.0322", "D28945-B-2", is_ring=True),
            ]),
            ("1405", False, [
                _item("F3.6061.0201", "D30001-A", is_ring=False, height_mm=703),
            ]),
        )
        template = Path(__file__).parents[1] / "legacy-planner" / "assets" / "锻造预热炉装炉图 v5.xlsx"

        exported = legacy_adapter.export_loading_map_with_openpyxl(result, template)
        imported = legacy_adapter.import_loading_map(exported)

        self.assertTrue(imported["ok"])
        self.assertEqual(imported["imported_items"], 3)
        imported_rows = legacy_adapter._export_furnace_details(imported["result"])
        self.assertIn(["1404", "F2.6061.0322", "D28945-B-1~2", 2, "430mm", None], imported_rows)
        self.assertIn(["1405", "F3.6061.0201", "D30001-A", 1, "703mm", 556.07], imported_rows)

        reexported = legacy_adapter.export_loading_map_with_openpyxl(imported["result"], template)
        sheet = openpyxl.load_workbook(BytesIO(reexported))["预热炉装炉图"]
        self.assertEqual(sheet["P10"].value, "D28945-B-1~2")
        self.assertEqual(sheet["S11"].value, 556.07)
        self.assertEqual(sheet["N10"].value, "1404")
        self.assertEqual(sheet["N11"].value, "1405")


class StandalonePlannerHttpTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.original_state_root = app.STATE_ROOT
        app.STATE_ROOT = Path(self.temp.name)
        self.server = app.create_server(port=0)
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        self.base_url = f"http://127.0.0.1:{self.server.server_port}"

    def tearDown(self):
        self.server.shutdown()
        self.server.server_close()
        self.thread.join(timeout=2)
        app.STATE_ROOT = self.original_state_root
        self.temp.cleanup()

    def request(self, path, body=None):
        data = None if body is None else json.dumps(body, ensure_ascii=False).encode("utf-8")
        headers = {} if data is None else {"Content-Type": "application/json"}
        try:
            response = urlopen(Request(f"{self.base_url}{path}", data=data, headers=headers))
        except HTTPError as exc:
            return exc.code, exc.headers, exc.read()
        return response.status, response.headers, response.read()

    def test_serves_independent_page_and_health(self):
        status, _, body = self.request("/")
        self.assertEqual(status, 200)
        page = body.decode("utf-8")
        self.assertIn("装炉排程（独立版）", page)
        self.assertNotIn("TIMELINE CAPACITY", page)
        self.assertNotIn("时间轴可排炉", page)
        self.assertNotIn("furnaceSummary", page)
        self.assertNotIn("furnaceSelect", page)
        self.assertNotIn("selectedTitle", page)
        self.assertNotIn("clearPlan", page)

        status, _, body = self.request("/api/v1/health")
        self.assertEqual(status, 200)
        health = json.loads(body)
        self.assertTrue(health["ok"])
        self.assertEqual(health["data_root"], str(app.STATE_ROOT))

    def test_legacy_routes_and_local_state_are_independent(self):
        status, _, body = self.request("/legacy-planner.html")
        self.assertEqual(status, 200)
        page = body.decode("utf-8")
        self.assertIn("const entityFurnaces = ['1412', '1413']", page)
        self.assertIn("导入排炉图", page)
        self.assertIn("/api/legacy/loading-plans/import", page)
        self.assertNotIn("legacyImportPanel", page)
        self.assertNotIn("导入到当前装炉安排", page)

        status, _, body = self.request("/api/legacy/billet-lengths")
        self.assertEqual(status, 200)
        self.assertIn("lengths", json.loads(body))

        # 独立版不写任何看板状态文件：既没有主看板数据库，也没有已删除的
        # v1 状态文件；装炉安排只存在于浏览器会话和导出的装炉图中。
        leftovers = {entry.name for entry in app.STATE_ROOT.iterdir()}
        self.assertNotIn("continuous_workbench.db", leftovers)
        self.assertNotIn("planner_state.json", leftovers)

    def test_removed_v1_state_routes_return_404(self):
        for path in ("/api/v1/dashboard", "/api/v1/loading/availability"):
            status, _, _ = self.request(path)
            self.assertEqual(status, 404)

    def test_loading_calculation_excludes_1412_and_1413(self):
        plan = DEFAULT_SERVICE.allocate(
            [{
                "code": "F3.6061.STANDALONE",
                "alloy": "锻造",
                "height_mm": 1200,
                "diameter_mm": 630,
                "rod_length_mm": 1200,
                "quantity": 1,
            }]
        )
        self.assertNotIn("1412", plan["allowed_furnace_ids"])
        self.assertNotIn("1413", plan["allowed_furnace_ids"])


if __name__ == "__main__":
    unittest.main()
