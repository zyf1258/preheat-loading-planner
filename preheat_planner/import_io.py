"""装炉图 Excel 导入与位置恢复。

导出图同时包含左侧可视炉位和右侧完整明细。导入时先读右侧明细建立完整工件
清单，再用左侧格位恢复已经显示的列位，最后把模板可视区域之外的工件放回原炉
的可用列。这样人工编辑后的装炉图仍可以继续换炉和再次导出。
"""

from __future__ import annotations

import re
from io import BytesIO
from typing import Any, Callable, Mapping

import openpyxl

from .export_io import _export_layouts
from .manual_service import (
    add_unconfigured_entity_furnaces,
    add_virtual_furnace,
    refresh_column,
)
from .rules import FURNACES, RING_HEIGHT, RING_LIMIT


def _import_text(value: object) -> str:
    """把 Excel 单元格统一转换成去除首尾空白的文本。"""

    return str(value or "").strip()


def _import_number(value: object) -> float:
    """读取普通数字或类似 ``430mm`` 的带单位文本。"""

    if isinstance(value, (int, float)):
        return float(value)
    match = re.search(r"-?\d+(?:\.\d+)?", _import_text(value).replace(",", ""))
    if not match:
        return 0.0
    try:
        return float(match.group(0))
    except ValueError:
        return 0.0


def expand_forging_numbers(value: object, quantity: int) -> list[str | None]:
    """将 ``D28945-B-1~2`` 展开为两个独立锻造号。"""

    count = max(int(quantity or 1), 1)
    text = _import_text(value).replace("～", "~")
    if not text:
        return [None] * count
    match = re.fullmatch(r"(.+-)(\d+)~(\d+)", text)
    if match:
        prefix, start, end = match.groups()
        start_number, end_number = int(start), int(end)
        if end_number >= start_number:
            expanded = [f"{prefix}{number}" for number in range(start_number, end_number + 1)]
            if len(expanded) >= count:
                return expanded[:count]
            return expanded + [None] * (count - len(expanded))
    return [text] * count


def _import_furnace_id(value: object) -> str:
    """标准化导出表中的炉号。"""

    text = _import_text(value).replace("－", "-")
    return re.sub(r"^YJ-", "", text, flags=re.IGNORECASE)


def read_detail_records(sheet: Any) -> list[dict]:
    """读取右侧装炉明细并恢复被合并单元格隐藏的炉号、物料号。"""

    required_headers = {"炉号", "物料号", "锻造号", "数量", "高度"}
    header_row = None
    header_columns = {}
    for row_number in range(1, min(sheet.max_row, 80) + 1):
        columns = {
            _import_text(sheet.cell(row_number, column_number).value): column_number
            for column_number in range(1, sheet.max_column + 1)
        }
        if required_headers.issubset(columns):
            header_row = row_number
            header_columns = columns
            break
    if header_row is None:
        raise ValueError("导入文件缺少“装炉明细”表头，请选择系统导出的装炉图")

    app_specs = {spec.furnace_id: spec for spec in FURNACES}
    records = []
    last_furnace = ""
    last_material = ""
    for row_number in range(header_row + 1, sheet.max_row + 1):
        raw_furnace = sheet.cell(row_number, header_columns["炉号"]).value
        raw_material = sheet.cell(row_number, header_columns["物料号"]).value
        raw_forging = sheet.cell(row_number, header_columns["锻造号"]).value
        raw_quantity = sheet.cell(row_number, header_columns["数量"]).value
        raw_height = sheet.cell(row_number, header_columns["高度"]).value
        raw_weight = (
            sheet.cell(row_number, header_columns["预估重量"]).value
            if "预估重量" in header_columns
            else None
        )
        if not any(
            _import_text(value)
            for value in (raw_furnace, raw_material, raw_forging, raw_quantity, raw_height, raw_weight)
        ):
            if records:
                break
            continue

        furnace_id = _import_furnace_id(raw_furnace) or last_furnace
        if _import_text(raw_furnace):
            if furnace_id != last_furnace:
                last_material = ""
            last_furnace = furnace_id
        material_code = _import_text(raw_material) or last_material
        if _import_text(raw_material):
            last_material = material_code
        if not furnace_id or material_code in {"无实际装炉工件", "无"}:
            continue
        if furnace_id not in app_specs:
            raise ValueError(f"装炉明细第{row_number}行包含未知炉号：{furnace_id}")

        quantity = max(int(_import_number(raw_quantity) or 1), 1)
        forging_no = _import_text(raw_forging) or None
        height_mm = _import_number(raw_height)
        estimated_raw = _import_text(raw_weight)
        estimated_weight = _import_number(raw_weight) if estimated_raw else None
        is_ring = bool(app_specs[furnace_id].is_ring)
        if not is_ring and height_mm == RING_HEIGHT and estimated_weight is None:
            is_ring = True
        if height_mm <= 0 and is_ring:
            height_mm = RING_HEIGHT
        records.append(
            {
                "furnace_id": furnace_id,
                "code": material_code,
                "forging_no": forging_no,
                "quantity": quantity,
                "height_mm": height_mm,
                "is_ring": is_ring,
            }
        )
    return records


def _item_from_record(record: dict, forging_no: str | None) -> dict:
    """把导入明细的一根工件转换成页面结果字典。"""

    code = record["code"]
    return {
        "code": code,
        "alloy": "",
        "quantity": 1,
        "height_mm": record["height_mm"],
        "rod_length_mm": record["height_mm"],
        "diameter_mm": 950 if "950" in code else 630,
        "is_ring": record["is_ring"],
        "forging_no": forging_no,
        "cumulative_mm": 0,
        "type": "环件" if record["is_ring"] else "导入装炉图",
    }


def _take_unit(index: dict, key: tuple[str, str]) -> dict | None:
    """按炉号和显示文本取出一件尚未恢复的明细。"""

    for unit in index.get(key, []):
        if not unit["used"]:
            unit["used"] = True
            return unit
    return None


def _append_unassigned(result: dict, item: dict, reason: str) -> None:
    """把无法恢复的工件放入未分配列表。"""

    result.setdefault("unassigned", []).append(
        {
            "code": item.get("code", ""),
            "forging_no": item.get("forging_no"),
            "reason": reason,
            "quantity": 1,
        }
    )


def _place_remaining(result: dict, furnace_id: str, item: dict) -> None:
    """将右侧明细中未出现在左侧格位的工件放回原炉。"""

    furnace = (result.get("furnaces") or {}).get(furnace_id)
    spec = (furnace or {}).get("spec") or {}
    height_mm = float(item.get("height_mm") or 0)
    if not furnace or spec.get("is_virtual") or height_mm <= 0:
        _append_unassigned(result, item, "导入后没有可用炉位")
        return
    limit = RING_LIMIT if item.get("is_ring") and spec.get("is_ring") else float(spec.get("height_mm") or 0)
    candidates = [
        (float(column.get("height_mm") or 0), index)
        for index, column in enumerate(furnace.get("columns") or [])
        if float(column.get("height_mm") or 0) + height_mm <= limit
    ]
    if not candidates:
        _append_unassigned(result, item, "导入后原炉没有满足高度约束的炉位")
        return
    _, column_index = min(candidates)
    column = furnace["columns"][column_index]
    column["items"].append(item)
    refresh_column(column)


def restore_loading_map(
    map_bytes: bytes,
    empty_result_factory: Callable[[], dict],
) -> dict:
    """从导出图恢复可继续编辑的排炉结果。"""

    if not map_bytes:
        raise ValueError("请选择要导入的装炉图")
    try:
        workbook = openpyxl.load_workbook(BytesIO(map_bytes), data_only=True)
    except Exception as exc:
        raise ValueError(f"无法读取装炉图 Excel：{exc}") from exc
    try:
        sheet = workbook["预热炉装炉图"] if "预热炉装炉图" in workbook.sheetnames else workbook.active
        records = read_detail_records(sheet)
        result = empty_result_factory()
        add_unconfigured_entity_furnaces(result)

        units = []
        lookup = {}
        for record in records:
            for forging_no in expand_forging_numbers(record["forging_no"], record["quantity"]):
                unit = {
                    "furnace_id": record["furnace_id"],
                    "item": _item_from_record(record, forging_no),
                    "used": False,
                }
                units.append(unit)
                furnace_id = record["furnace_id"]
                lookup.setdefault((furnace_id, record["code"]), []).append(unit)
                if forging_no:
                    lookup.setdefault((furnace_id, forging_no), []).append(unit)

        for furnace_id, layout in _export_layouts().items():
            furnace = result["furnaces"].get(furnace_id)
            if not furnace:
                continue
            for column_index, cell_addresses in enumerate(layout["cells"]):
                if column_index >= len(furnace.get("columns") or []):
                    continue
                column = furnace["columns"][column_index]
                for address in reversed(cell_addresses):
                    label = _import_text(sheet[address].value)
                    if not label:
                        continue
                    unit = _take_unit(lookup, (furnace_id, label))
                    if unit is None:
                        raise ValueError(f"{furnace_id} 炉位 {address} 的“{label}”无法在右侧明细中匹配")
                    column["items"].append(unit["item"])
                refresh_column(column)

        for unit in units:
            if not unit["used"]:
                _place_remaining(result, unit["furnace_id"], unit["item"])

        result["total_requested"] = len(units)
        add_virtual_furnace(result, [unit["item"] for unit in units])
        return {
            "ok": True,
            "items": [unit["item"] for unit in units],
            "result": result,
            "imported_items": len(units),
        }
    finally:
        workbook.close()


__all__ = ["expand_forging_numbers", "read_detail_records", "restore_loading_map"]
