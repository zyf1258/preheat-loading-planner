"""排炉相关 Excel 输入读取器。

本模块只负责把 Excel 文件转换成普通 Python 数据，不执行排炉，也不负责保存
网页状态。Excel 的列位置和中文表头属于文件格式问题，集中放在这里后，业务层
不需要再夹杂 openpyxl 细节。
"""

from __future__ import annotations

from collections import Counter
from collections.abc import Callable, Mapping
from pathlib import Path

import openpyxl


def read_billet_length_map(path: Path, number_parser: Callable[[object], float | int]) -> dict:
    """读取 MES 棒料登记表，返回每个物料编码的常用长度。

    同一物料编码出现多个长度时取出现次数最多的长度，保持原页面的记忆规则。
    ``number_parser`` 由兼容层传入，便于沿用旧项目对 Excel 数字的解释方式。
    """

    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            first_row = next(rows)
        except StopIteration as exc:
            raise ValueError("棒料登记表为空") from exc
        headers = {
            str(value or "").strip(): index
            for index, value in enumerate(first_row, start=1)
        }
        required = {"产品编号", "长度"}
        if not required.issubset(headers):
            raise ValueError("棒料登记表缺少“产品编号”或“长度”列")
        code_index = headers["产品编号"] - 1
        length_index = headers["长度"] - 1
        values_by_code = {}
        for row in rows:
            code = str(row[code_index] or "").strip() if code_index < len(row) else ""
            length = number_parser(row[length_index]) if length_index < len(row) else 0
            if code and length > 0:
                values_by_code.setdefault(code, []).append(length)
        return {code: Counter(values).most_common(1)[0][0] for code, values in values_by_code.items()}
    finally:
        workbook.close()


def load_plan_items(
    path: Path,
    lengths: Mapping[str, object],
    *,
    ring_height: int,
    number_parser: Callable[[object], float | int],
    forging_parser: Callable[[object, int], list[str] | None],
) -> tuple[list[dict], list[dict]]:
    """读取生产计划表并拆分为可排炉工件和缺长度记录。

    生产计划沿用当前文件格式：B 列物料编码、C 列工艺描述、D 列锻造号、E 列
    数量。锻造号能够按数量完整展开时，每根工件保留独立锻造号；否则保留整行
    数量，和原页面行为一致。
    """

    workbook = openpyxl.load_workbook(path, data_only=True)
    try:
        sheet = workbook.active
        items = []
        for row_number in range(2, sheet.max_row + 1):
            code = str(sheet[f"B{row_number}"].value or "").strip()
            alloy = str(sheet[f"C{row_number}"].value or "").strip()
            quantity = number_parser(sheet[f"E{row_number}"].value)
            forging_numbers = forging_parser(sheet[f"D{row_number}"].value, quantity)
            if not code or quantity <= 0 or "外发" in alloy or "下料" in alloy:
                continue
            is_ring = any(token in alloy for token in ("锻造碾环", "碾环", "冲孔"))
            height = ring_height if is_ring else lengths.get(code, 0)
            if height <= 0:
                items.append(
                    {
                        "code": code,
                        "alloy": alloy,
                        "quantity": quantity,
                        "height_mm": 0,
                        "diameter_mm": 630,
                        "is_ring": False,
                        "invalid": "未记住该棒料长度",
                    }
                )
                continue
            base = {
                "code": code,
                "alloy": alloy,
                "height_mm": height,
                "diameter_mm": 950 if "950" in alloy else 630,
                "is_ring": is_ring,
            }
            if forging_numbers:
                items.extend(
                    {**base, "quantity": 1, "forging_no": forging_no}
                    for forging_no in forging_numbers
                )
            else:
                items.append({**base, "quantity": quantity})
        missing = [item for item in items if item.get("invalid")]
        valid = [item for item in items if not item.get("invalid")]
        return valid, [
            {
                "code": item["code"],
                "forging_no": item.get("forging_no"),
                "reason": item["invalid"],
                "quantity": item["quantity"],
            }
            for item in missing
        ]
    finally:
        workbook.close()


__all__ = ["load_plan_items", "read_billet_length_map"]
