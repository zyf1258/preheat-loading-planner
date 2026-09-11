"""装炉图 Excel 导出层。

本模块负责把排炉结果写入现有的《预热炉装炉图》模板。它只关心 Excel 版式、
明细分组和显示计算，不负责读取生产计划，也不负责执行排炉。

导出规则集中在这里：

* 环件炉中的环件按相同物料编码和高度合并；
* 同一前缀下连续锻造号压缩成 ``1~2`` 形式；
* 普通锻造炉不合并物料；
* 炉号列始终逐行写入，不纵向合并；
* 环件预估重量为空，其他工件按 ``长度 * 0.791`` 计算。
"""

from __future__ import annotations

import copy
import re
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from .service import DEFAULT_SERVICE


# 左侧炉子标题沿用原项目的总重量估算规则；右侧明细使用单独的业务规则。
EXPORT_WEIGHT_DIVISOR = 1.264
EXPORT_ESTIMATED_WEIGHT_PER_MM = 0.791


def _export_column_groups(columns: list[str], groups: list[list[int]]) -> list[list[str]]:
    """根据模板列字母和行号组装逻辑炉列的单元格地址。"""

    return [[f"{column}{row}" for row in rows] for rows in groups for column in columns]


def _export_write_row(sheet: Any, row_number: int, values: list[object]) -> None:
    """从 A 列开始把一行普通 Python 值写入工作表。"""

    for column_number, value in enumerate(values, start=1):
        sheet.cell(row=row_number, column=column_number).value = value


def _export_layouts() -> dict[str, dict[str, object]]:
    """返回现有 Excel 模板中各台配置炉的固定格位布局。"""

    return {
        "1408": {
            "title": "A4",
            "cells": [["A6", "A7", "A8"], ["B6", "B7", "B8"], ["C6", "C7", "C8"], ["A10", "A11", "A12"]],
        },
        "1405": {"title": "D4", "cells": _export_column_groups(["D", "E", "F"], [[6, 7, 8], [10, 11, 12]])},
        "1409": {"title": "G4", "cells": _export_column_groups(["G", "H", "I"], [[6, 7, 8], [10, 11, 12]])},
        "1404": {"title": "J4", "cells": _export_column_groups(["J", "K", "L"], [[6, 7, 8], [10, 11, 12]])},
        "1403": {"title": "A13", "cells": _export_column_groups(["A", "B", "C"], [[15, 16, 17], [19, 20, 21]])},
        "1406": {"title": "D13", "cells": _export_column_groups(["D", "E", "F"], [[15, 16, 17], [19, 20, 21]])},
        "1407": {"title": "H13", "cells": _export_column_groups(["H", "I", "J", "K"], [[15, 16, 17], [19, 20, 21], [23, 24, 25]])},
        "1410": {"title": "A22", "cells": _export_column_groups(["A", "B", "C"], [[25, 26, 27], [29, 30, 31]])},
        "1411": {"title": "D22", "cells": _export_column_groups(["D", "E", "F"], [[25, 26, 27], [29, 30, 31]])},
    }


def _export_fill(hex_color: str):
    """创建 Excel 实心填充。"""

    from openpyxl.styles import PatternFill

    return PatternFill(fill_type="solid", fgColor=hex_color.lstrip("#"))


def _export_border():
    """创建导出表格共用的细边框。"""

    from openpyxl.styles import Border, Side

    side = Side(style="thin", color="777777")
    return Border(left=side, right=side, top=side, bottom=side)


def _export_title_border():
    """创建炉号标题和时间框使用的中等边框。"""

    from openpyxl.styles import Border, Side

    side = Side(style="medium", color="777777")
    return Border(left=side, right=side, top=side, bottom=side)


def _export_merged_range_for_cell(sheet: Any, address: str):
    """找到包含标题单元格的合并区域。"""

    cell = sheet[address]
    for merged_range in sheet.merged_cells.ranges:
        if (
            merged_range.min_row <= cell.row <= merged_range.max_row
            and merged_range.min_col <= cell.column <= merged_range.max_col
        ):
            return merged_range
    return None


def _export_prepare_time_panels(sheet: Any, layouts: dict[str, dict[str, object]]) -> None:
    """把每个炉号标题拆成“进炉时间｜炉号｜出炉时间”三块。

    模板原本把整行标题合并为一个单元格。为了不改变炉位网格和右侧明细的
    坐标，这里只拆分标题区域本身：两侧保留可手工填写的空框，中间继续
    写入炉号、角色、重量和预热时间。时间数据当前没有来自排程结果的来源，
    因此导出时故意留空，方便现场在 Excel 中补填。
    """

    from openpyxl.styles import Alignment

    time_border = _export_title_border()
    for layout in layouts.values():
        title_address = str(layout["title"])
        merged_range = _export_merged_range_for_cell(sheet, title_address)
        if merged_range is None:
            continue

        min_row, max_row = merged_range.min_row, merged_range.max_row
        min_col, max_col = merged_range.min_col, merged_range.max_col
        source = sheet[title_address]
        source_style = copy.copy(source._style)
        source_fill = copy.copy(source.fill)
        source_font = copy.copy(source.font)
        sheet.unmerge_cells(str(merged_range))

        # 先把原来的标题样式铺回整个区域，再分别覆盖两侧时间框样式。
        for row in range(min_row, max_row + 1):
            for column in range(min_col, max_col + 1):
                cell = sheet.cell(row=row, column=column)
                cell.value = None
                cell._style = copy.copy(source_style)
                cell.border = time_border

        left_range = f"{sheet.cell(min_row, min_col).coordinate}:{sheet.cell(max_row, min_col).coordinate}"
        right_range = f"{sheet.cell(min_row, max_col).coordinate}:{sheet.cell(max_row, max_col).coordinate}"
        center_start_col = min_col + 1
        center_end_col = max_col - 1

        # 多行标题（1410、1411）需要把时间框纵向合并；普通标题只占一行。
        if max_row > min_row:
            sheet.merge_cells(left_range)
            sheet.merge_cells(right_range)
        if center_start_col <= center_end_col:
            center_range = (
                f"{sheet.cell(min_row, center_start_col).coordinate}:"
                f"{sheet.cell(max_row, center_end_col).coordinate}"
            )
            if center_end_col > center_start_col or max_row > min_row:
                sheet.merge_cells(center_range)
        else:
            center_range = None

        left = sheet.cell(min_row, min_col)
        right = sheet.cell(min_row, max_col)
        side_font = copy.copy(source_font)
        side_font.sz = min(float(side_font.sz or 11), 9)
        side_font.bold = True
        left.value = "进炉时间\n\n"
        right.value = "出炉时间\n\n"
        for side in (left, right):
            side.fill = _export_fill("FFFFFF")
            side.font = copy.copy(side_font)
            side.alignment = Alignment(
                wrap_text=True,
                horizontal="center",
                vertical="top",
            )
            side.border = time_border

        # 让中间标题仍然使用原来的浅蓝底色和居中样式。
        center = sheet.cell(min_row, center_start_col) if center_range else source
        center.fill = copy.copy(source_fill)
        center.font = copy.copy(source_font)
        center.alignment = Alignment(
            wrap_text=True,
            horizontal="center",
            vertical="center",
        )
        center.border = time_border
        if max_row > min_row:
            sheet.row_dimensions[min_row].height = max(
                sheet.row_dimensions[min_row].height or 15,
                75,
            )
            for row in range(min_row + 1, max_row + 1):
                sheet.row_dimensions[row].height = max(
                    sheet.row_dimensions[row].height or 15,
                    18,
                )
        layout["title"] = center.coordinate


def _export_item_text(item: dict) -> str:
    """取得左侧炉位格中显示的锻造号或物料编码。"""

    return item.get("forging_no") or item.get("code") or "未命名工件"


def _export_item_weight_kg(item: dict) -> float | None:
    """获取左侧炉子标题和预热时间使用的单件重量。"""

    for key in ("weight_kg", "weight", "重量", "重量kg"):
        try:
            value = float(item.get(key))
        except (TypeError, ValueError):
            continue
        if value > 0:
            return value
    for key in ("rod_length_mm", "rod_length", "length_mm", "height_mm"):
        try:
            length = float(item.get(key))
        except (TypeError, ValueError):
            continue
        if length > 0:
            return length / EXPORT_WEIGHT_DIVISOR
    return None


def _export_furnace_items(furnace: dict) -> list[dict]:
    """把炉内的二维列结构摊平成明细需要的一维工件列表。"""

    return [
        item
        for column in (furnace.get("columns") or [])
        for item in (column.get("items") or [])
        if isinstance(item, dict)
    ]


def _export_item_quantity(item: dict) -> float:
    """读取数量，并将缺失、非法或非正值兜底为 1。"""

    try:
        quantity = float(item.get("quantity") or 1)
    except (TypeError, ValueError):
        quantity = 1.0
    return quantity if quantity > 0 else 1.0


def _export_item_height(item: dict) -> float | None:
    """从兼容字段中取得工件长度/高度（mm）。"""

    for key in ("height_mm", "rod_length_mm", "rod_length", "length_mm"):
        try:
            height = float(item.get(key))
        except (TypeError, ValueError):
            continue
        if height > 0:
            return height
    return None


def _export_item_estimated_weight(item: dict) -> float | None:
    """计算右侧明细的预估重量，环件不计算。"""

    if item.get("is_ring"):
        return None
    height = _export_item_height(item)
    return None if height is None else round(height * EXPORT_ESTIMATED_WEIGHT_PER_MM, 2)


def _export_furnace_sort_key(furnace_id: object):
    """让数字炉号按数字顺序排序，同时兼容虚拟炉等文本编号。"""

    try:
        return (0, int(furnace_id))
    except (TypeError, ValueError):
        return (1, str(furnace_id))


def _export_text_sort_key(value: object):
    """按数字片段的数值排序，避免 D10 排在 D2 前面。"""

    text = str(value or "")
    return tuple(
        (0, int(part)) if part.isdigit() else (1, part.lower())
        for part in re.split(r"(\d+)", text)
    )


def _export_detail_value(value: float) -> int | float:
    """将整数形式的浮点数量转成整数，改善 Excel 显示。"""

    return int(value) if value.is_integer() else value


def _export_detail_height_text(height_key: object) -> str:
    """将高度键转换为带 mm 单位的文本。"""

    return f"{height_key:g}mm" if isinstance(height_key, float) else str(height_key)


def _export_numeric_forging_no(value: object):
    """拆出可合并锻造号的前缀和末尾数字。"""

    text = str(value or "").strip()
    match = re.fullmatch(r"(.+-)(\d+)", text)
    return (match.group(1), int(match.group(2))) if match else None


def _export_ring_forging_groups(entries: list[tuple[str, float]]) -> list[tuple[str, float]]:
    """合并同一物料下连续的锻造号，并保留每段的实际数量。"""

    numeric = {}
    plain = {}
    for forging_no, quantity in entries:
        parsed = _export_numeric_forging_no(forging_no)
        if parsed is None:
            plain[forging_no] = plain.get(forging_no, 0.0) + quantity
        else:
            numeric[parsed] = numeric.get(parsed, 0.0) + quantity

    collapsed = []
    for prefix in sorted({prefix for prefix, _ in numeric}, key=_export_text_sort_key):
        values = sorted(number for current_prefix, number in numeric if current_prefix == prefix)
        start = previous = values[0]
        for number in values[1:] + [None]:
            if number is not None and number == previous + 1:
                previous = number
                continue
            forging_no = (
                f"{prefix}{start}"
                if start == previous
                else f"{prefix}{start}~{previous}"
            )
            quantity = sum(numeric[(prefix, current)] for current in range(start, previous + 1))
            collapsed.append((forging_no, quantity))
            if number is not None:
                start = previous = number
    collapsed.extend(sorted(plain.items(), key=lambda entry: _export_text_sort_key(entry[0])))
    return sorted(collapsed, key=lambda entry: _export_text_sort_key(entry[0]))


def _export_furnace_detail_groups(result: dict) -> list[dict]:
    """生成右侧明细行，并标记哪些物料号可以纵向合并。"""

    from collections import defaultdict

    detail_groups = []
    furnaces = result.get("furnaces") or {}
    for furnace_id in sorted(furnaces, key=_export_furnace_sort_key):
        furnace = furnaces[furnace_id]
        if not isinstance(furnace, dict):
            continue
        spec = furnace.get("spec") or {}
        if spec.get("is_virtual") or str(furnace_id).upper().startswith("VIRTUAL"):
            continue

        ring_furnace = bool(spec.get("is_ring"))
        ring_by_material = defaultdict(list)
        exact_groups = defaultdict(float)
        for item in _export_furnace_items(furnace):
            material_code = str(item.get("code") or "").strip() or "未命名工件"
            forging_no = str(item.get("forging_no") or "").strip() or material_code
            height = _export_item_height(item)
            height_key = height if height is not None else "待补充"
            estimated_weight = _export_item_estimated_weight(item)
            quantity = _export_item_quantity(item)
            if ring_furnace and item.get("is_ring"):
                ring_by_material[(material_code, height_key)].append((forging_no, quantity))
            else:
                exact_groups[(material_code, forging_no, height_key, estimated_weight)] += quantity

        for (material_code, height_key), entries in ring_by_material.items():
            for forging_no, quantity in _export_ring_forging_groups(entries):
                detail_groups.append(
                    {
                        "values": [
                            str(furnace_id),
                            material_code,
                            forging_no,
                            _export_detail_value(quantity),
                            _export_detail_height_text(height_key),
                            None,
                        ],
                        "merge_furnace": ring_furnace,
                        "merge_material": True,
                        "material_key": (str(furnace_id), material_code),
                    }
                )

        for (material_code, forging_no, height_key, estimated_weight), quantity in exact_groups.items():
            detail_groups.append(
                {
                    "values": [
                        str(furnace_id),
                        material_code,
                        forging_no,
                        _export_detail_value(quantity),
                        _export_detail_height_text(height_key),
                        estimated_weight,
                    ],
                    "merge_furnace": ring_furnace,
                    "merge_material": False,
                    "material_key": None,
                }
            )

    detail_groups.sort(
        key=lambda group: (
            _export_furnace_sort_key(group["values"][0]),
            _export_text_sort_key(group["values"][1]),
            0 if group["merge_material"] else 1,
            _export_text_sort_key(group["values"][2]),
            _export_text_sort_key(group["values"][4]),
        )
    )
    return detail_groups


def export_furnace_details(result: dict) -> list[list[object]]:
    """只生成右侧明细值，便于测试和其他导出格式复用。"""

    return [group["values"] for group in _export_furnace_detail_groups(result)]


def _export_furnace_total_weight(furnace: dict) -> float | None:
    """计算左侧炉子标题显示用的完整总重量。"""

    total = 0.0
    for item in _export_furnace_items(furnace):
        weight = _export_item_weight_kg(item)
        if weight is None:
            return None
        quantity = _export_item_quantity(item)
        total += weight * quantity
    return total


def _export_furnace_title(furnace_id: str, furnace: dict) -> str:
    """生成左侧炉号、角色、总重量和预计预热时间标题。"""

    total_weight = _export_furnace_total_weight(furnace)
    if total_weight is None:
        weight_text = "待补充"
        time_text = "待补充"
    elif total_weight <= 0:
        weight_text = "0.0 kg"
        time_text = "--"
    else:
        weight_text = f"{total_weight:.1f} kg"
        prediction = DEFAULT_SERVICE.calculate_preheat_time(furnace_id, total_weight)
        predicted_hours = prediction.get("predicted_hours")
        if predicted_hours is not None:
            time_text = f"{float(predicted_hours):.2f} 小时"
        elif furnace_id in {"1406", "1407"}:
            time_text = "按现有逻辑"
        else:
            time_text = "待手动确认"
    role = (furnace.get("spec") or {}).get("role", "")
    return f"{furnace_id}　{role}\n总重量：{weight_text}　预计预热时间：{time_text}"


def export_loading_map_with_openpyxl(result: dict, template_path: Path) -> bytes:
    """把排炉结果写入 Excel 模板并返回可下载的二进制内容。"""

    import openpyxl
    from openpyxl.styles import Alignment, Font

    workbook = openpyxl.load_workbook(template_path)
    try:
        sheet = workbook["预热炉装炉图"] if "预热炉装炉图" in workbook.sheetnames else workbook.active
        sheet["A3"] = f"导出时间：{datetime.now().strftime('%Y/%m/%d %H:%M:%S')}"
        sheet["G3"] = f"计划数量：{result.get('total_requested', 0)}　已装入：{result.get('total_placed', 0)}"
        layouts = _export_layouts()
        _export_prepare_time_panels(sheet, layouts)
        grid_overflow = []
        item_border = _export_border()

        for column_letter in "ABCDEFGHIJKL":
            sheet.column_dimensions[column_letter].width = 18

        for furnace_id, layout in layouts.items():
            cells = layout["cells"]
            for cell_addresses in cells:
                for address in cell_addresses:
                    sheet[address].value = None
            furnace = (result.get("furnaces") or {}).get(furnace_id)
            if not furnace:
                continue
            title_cell = sheet[layout["title"]]
            title_cell.value = _export_furnace_title(furnace_id, furnace)
            title_cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
            sheet.row_dimensions[title_cell.row].height = max(sheet.row_dimensions[title_cell.row].height or 15, 60)
            for column_index, cell_addresses in enumerate(cells):
                columns = furnace.get("columns") or []
                items = columns[column_index].get("items", []) if column_index < len(columns) else []
                for item_index, item in enumerate(items[: len(cell_addresses)]):
                    target = sheet[cell_addresses[len(cell_addresses) - 1 - item_index]]
                    target.value = _export_item_text(item)
                    target.fill = _export_fill("FDE9B4" if item.get("is_ring") else "DDEEF4")
                    target.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
                    target.border = item_border
                    item_font = copy.copy(target.font)
                    item_font.sz = 14
                    target.font = item_font
                for item_index, item in enumerate(items[len(cell_addresses) :], start=len(cell_addresses) + 1):
                    grid_overflow.append(
                        {
                            **item,
                            "furnace_id": furnace_id,
                            "column_index": column_index,
                            "layer_index": item_index,
                        }
                    )

        for row in sheet["A38:C400"]:
            for cell in row:
                cell.value = None
        unassigned_items = list(
            (((result.get("furnaces") or {}).get("VIRTUAL") or {}).get("columns") or [{}])[0].get("items", [])
            or result.get("unassigned")
            or []
        )
        unassigned_start_row = 38
        if grid_overflow:
            sheet[f"A{unassigned_start_row}"] = f"第四层备注（{len(grid_overflow)} 根）"
            sheet[f"A{unassigned_start_row}"].fill = _export_fill("FFF2CC")
            sheet[f"A{unassigned_start_row}"].font = Font(bold=True)
            overflow_header = unassigned_start_row + 1
            _export_write_row(sheet, overflow_header, ["炉号", "列", "备注"])
            overflow_rows = []
            for item in grid_overflow:
                label = (
                    f"{item.get('forging_no')}（{item.get('code', '')}）"
                    if item.get("forging_no")
                    else item.get("code") or "未命名工件"
                )
                length_text = f"{item.get('height_mm')}mm" if item.get("height_mm") else "待补充长度"
                overflow_rows.append(
                    [
                        item["furnace_id"],
                        f"第{item['column_index'] + 1}列",
                        f"第{item['layer_index']}层：{label} / {length_text}",
                    ]
                )
            first = overflow_header + 1
            for row_offset, values in enumerate(overflow_rows):
                _export_write_row(sheet, first + row_offset, values)
            unassigned_start_row = first + len(overflow_rows) + 1

        sheet[f"A{unassigned_start_row}"] = f"未排炉工件（{len(unassigned_items)} 根）"
        sheet[f"A{unassigned_start_row}"].fill = _export_fill("FCE4D6")
        sheet[f"A{unassigned_start_row}"].font = Font(bold=True)
        header_row = unassigned_start_row + 1
        _export_write_row(sheet, header_row, ["物料编码", "锻造号", "棒料长度"])
        rows = [
            [
                item.get("code", ""),
                item.get("forging_no") or item.get("code", ""),
                f"{item.get('height_mm')}mm" if item.get("height_mm") else "待补充",
            ]
            for item in unassigned_items
        ] or [["无", "", ""]]
        first = header_row + 1
        for row_offset, values in enumerate(rows):
            _export_write_row(sheet, first + row_offset, values)

        detail_title_row = 8
        detail_header_row = detail_title_row + 1
        detail_first_row = detail_header_row + 1
        detail_groups = _export_furnace_detail_groups(result)
        detail_last_row = max(detail_first_row, detail_first_row + len(detail_groups) - 1)
        for merged_range in list(sheet.merged_cells.ranges):
            if merged_range.min_col >= 14 and merged_range.max_col <= 19 and merged_range.max_row >= detail_title_row:
                sheet.unmerge_cells(str(merged_range))
        for row in sheet.iter_rows(
            min_row=detail_title_row,
            max_row=max(detail_last_row, 400),
            min_col=14,
            max_col=19,
        ):
            for cell in row:
                cell.value = None
                cell.fill = copy.copy(sheet["N8"].fill)
                cell.border = copy.copy(sheet["N8"].border)
        sheet.merge_cells("N8:S8")
        detail_title = sheet["N8"]
        detail_title.value = "装炉明细（实际装炉）"
        detail_title.fill = _export_fill("D9EAF7")
        detail_title.font = Font(bold=True, size=12)
        detail_title.alignment = Alignment(horizontal="center", vertical="center")
        detail_title.border = item_border
        detail_headers = ["炉号", "物料号", "锻造号", "数量", "高度", "预估重量"]
        for column_number, value in enumerate(detail_headers, start=14):
            cell = sheet.cell(row=detail_header_row, column=column_number)
            cell.value = value
            cell.fill = _export_fill("5B9BD5")
            cell.font = Font(bold=True, color="FFFFFF", size=11)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = item_border
        detail_rows = [group["values"] for group in detail_groups] or [["", "无实际装炉工件", "", "", "", ""]]
        for row_offset, values in enumerate(detail_rows):
            row_number = detail_first_row + row_offset
            for column_number, value in enumerate(values, start=14):
                cell = sheet.cell(row=row_number, column=column_number)
                cell.value = value
                cell.font = Font(size=11)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = item_border
                if column_number == 19 and isinstance(value, (int, float)):
                    cell.number_format = "0.00"

        # 炉号列永不合并；只有 detail_groups 明确允许的环件物料号才纵向合并。
        row_index = 0
        while row_index < len(detail_groups):
            group = detail_groups[row_index]
            if not group["merge_material"]:
                row_index += 1
                continue
            start_index = row_index
            run_value = group["values"][1]
            row_index += 1
            while (
                row_index < len(detail_groups)
                and detail_groups[row_index]["merge_material"]
                and detail_groups[row_index]["values"][1] == run_value
            ):
                row_index += 1
            if row_index - start_index > 1:
                sheet.merge_cells(
                    start_row=detail_first_row + start_index,
                    start_column=15,
                    end_row=detail_first_row + row_index - 1,
                    end_column=15,
                )
                sheet.cell(detail_first_row + start_index, 15).alignment = Alignment(
                    horizontal="center", vertical="center"
                )

        for column, width in (("N", 12), ("O", 25), ("P", 25), ("Q", 10), ("R", 12), ("S", 14)):
            sheet.column_dimensions[column].width = width

        output = BytesIO()
        workbook.save(output)
        return output.getvalue()
    finally:
        workbook.close()


__all__ = [
    "EXPORT_ESTIMATED_WEIGHT_PER_MM",
    "EXPORT_WEIGHT_DIVISOR",
    "export_furnace_details",
    "export_loading_map_with_openpyxl",
]
